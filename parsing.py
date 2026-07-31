"""Stage 1 (Drashti): turn a real PDF/DOCX contract into a `Document`.

No hardcoded rules. The work is done by:
  - pdfplumber / python-docx  -> text extraction (with page info)
  - langdetect                -> source language
  - an LLM (Ollama by default)-> clause segmentation + translation

Everything else here is mechanical glue (looping, locating text back in the
original to compute char spans). Swap the LLM backend in `_llm_json` if you
use something other than Ollama.
"""
from __future__ import annotations
import os
import json
from uuid import uuid4
import requests
from langdetect import detect
from interfaces import Document, Clause

# ---- LLM backend (swap here if not using Ollama) ----
OLLAMA_URL = "http://localhost:11434/api/generate"
LLM_MODEL = "qwen2.5:3b"   # or "llama3.1", etc. — whatever you pulled in Ollama


def _llm_json(prompt: str) -> dict:
    """Send a prompt, get JSON back. Ollama's format=json forces valid JSON."""
    resp = requests.post(
        OLLAMA_URL,
        json={"model": LLM_MODEL, "prompt": prompt, "stream": False, "format": "json"},
        timeout=600,
    )
    resp.raise_for_status()
    return json.loads(resp.json()["response"])


# ---- Text extraction ----
def _extract_pages(path: str) -> list[tuple[int, str]]:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return [(i + 1, (pg.extract_text() or "")) for i, pg in enumerate(pdf.pages)]
    if ext == "docx":
        import docx
        d = docx.Document(path)
        return [(1, "\n".join(p.text for p in d.paragraphs))]
    raise ValueError(f"Unsupported file type '.{ext}'. Convert .doc to .docx first.")


def _build_fulltext(pages: list[tuple[int, str]]) -> tuple[str, list[tuple[int, int]]]:
    """Concatenate pages, remembering where each page starts (for char->page mapping)."""
    full, offsets = "", []
    for page_num, text in pages:
        offsets.append((len(full), page_num))
        full += text + "\n\n"
    return full, offsets


def _page_for_offset(offset: int, offsets: list[tuple[int, int]]) -> int:
    page = offsets[0][1] if offsets else 1
    for start, p in offsets:
        if offset >= start:
            page = p
        else:
            break
    return page


# ---- LLM-driven segmentation + translation (no rules) ----
_SEGMENT_PROMPT = """You are a legal contract parser. Split the contract below into its \
individual clauses. Return ONLY JSON of the form:
{"clauses": [{"heading": "<section heading, or null>", "text": "<verbatim clause text>"}]}
Copy each clause's text VERBATIM from the contract so it can be located again. \
Do not summarise, merge, or rewrite.

CONTRACT:
\"\"\"
%s
\"\"\""""


def _segment(full_text: str) -> list[dict]:
    return _llm_json(_SEGMENT_PROMPT % full_text).get("clauses", [])


def _translate_to_en(text: str) -> str:
    prompt = (
        'Translate the following contract clause into English. Preserve legal meaning. '
        'Return ONLY JSON: {"text_en": "..."}.\n\nCLAUSE:\n"""\n%s\n"""' % text
    )
    return _llm_json(prompt).get("text_en", text)


def _detect_lang(text: str) -> str:
    try:
        return detect(text)
    except Exception:
        return "en"


# ---- The stage-1 entry point (drop-in replacement for the stub) ----
def parse_document(file_path: str) -> Document:
    pages = _extract_pages(file_path)
    full_text, page_offsets = _build_fulltext(pages)
    lang = _detect_lang(full_text[:2000])
    raw_clauses = _segment(full_text)

    clauses: list[Clause] = []
    for i, rc in enumerate(raw_clauses):
        text_original = (rc.get("text") or "").strip()
        if not text_original:
            continue
        pos = full_text.find(text_original)
        start = pos if pos >= 0 else 0
        heading = rc.get("heading")
        if isinstance(heading, str) and heading.strip().lower() in ("", "null", "none"):
            heading = None
        clauses.append(Clause(
            clause_id=f"c_{i + 1:03d}",
            index=i,
            heading=heading,
            text_original=text_original,
            text_en=text_original if lang == "en" else _translate_to_en(text_original),
            page=_page_for_offset(start, page_offsets),
            char_span=(start, start + len(text_original)),
        ))

    return Document(
        doc_id="d_" + uuid4().hex[:10],
        filename=os.path.basename(file_path),
        source_language=lang,
        page_count=len(pages),
        clauses=clauses,
    )
